#!/usr/bin/env python3
"""
FSA Attendance Tracker
Load your weekly Excel sheets and explore attendance by week / month / quarter / year.
"""

import sys
import os
import re
from datetime import datetime

import pandas as pd
import openpyxl
import matplotlib
matplotlib.use("Qt5Agg")
from matplotlib.backends.backend_qt5agg import FigureCanvasQTAgg as FigureCanvas
from matplotlib.figure import Figure

from PyQt5.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QPushButton, QLabel, QComboBox, QTableWidget, QTableWidgetItem,
    QFileDialog, QTabWidget, QFrame, QSplitter, QHeaderView,
    QGroupBox, QDateEdit, QProgressDialog, QMessageBox, QLineEdit,
)
from PyQt5.QtCore import Qt, QDate, QThread, pyqtSignal
from PyQt5.QtGui import QColor, QFont

# ── Palette ───────────────────────────────────────────────────────────────────
BG        = "#1e1e2e"
BG_CARD   = "#2a2a3e"
BG_SIDE   = "#16162a"
BORDER    = "#334155"
TXT       = "#e2e8f0"
TXT2      = "#94a3b8"
BLUE      = "#4f9cf9"
GREEN     = "#4ade80"
RED       = "#f87171"
YELLOW    = "#fbbf24"
ORANGE    = "#fb923c"
PURPLE    = "#a78bfa"
TEAL      = "#2dd4bf"

STAT_COLORS = {
    "present":    GREEN,
    "absent":     RED,
    "lates":      YELLOW,
    "left_early": ORANGE,
    "sick":       PURPLE,
    "leave":      TEAL,
}

# ── Excel parsing ─────────────────────────────────────────────────────────────
SKIP = {"Master Attendance File", "Holiday Attendance", "Sheet1", "Sheet2"}

_DATE_FMTS = (
    "%m-%d-%Y", "%m-%d-%y", "%m/%d/%Y", "%m/%d/%y",
    "%m.%d.%Y", "%m.%d.%y",
)

def _sheet_date(sheet_name):
    """Try to extract a date from a sheet name like 'WE 7-2-23' or '7-14-24'."""
    s = re.sub(r"^WE\s*", "", sheet_name.strip(), flags=re.IGNORECASE).strip()
    s = re.sub(r"\s+", "-", s)
    for fmt in _DATE_FMTS:
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            pass
    return None


def _safe_int(v):
    if v is None:
        return 0
    try:
        return int(float(v))
    except (ValueError, TypeError):
        return 0


def load_excel(path, progress_cb=None):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheets = wb.sheetnames
    records = []

    for idx, name in enumerate(sheets):
        if progress_cb:
            progress_cb(idx, len(sheets), name)
        if name in SKIP:
            continue

        ws = wb[name]

        all_rows = list(ws.iter_rows(values_only=True))
        if len(all_rows) < 2:
            continue

        header = list(all_rows[0])
        if not header or header[0] != "Week Ending":
            continue

        # Required columns
        try:
            ci_name = header.index("NAME")
            ci_acct = header.index("Account")
            ci_sup  = header.index("Supervisor")
            ci_late = header.index("Lates")
            ci_le   = header.index("Left Early")
            ci_abs  = header.index("Absent")
            ci_sick = header.index("Sick")
            ci_lv   = header.index("Leave")
        except ValueError:
            continue

        ci_pres = header.index("Present") if "Present" in header else None
        ci_pct  = header.index("Weekly Attendance %") if "Weekly Attendance %" in header else None
        day_cols = list(range(5, ci_late))  # Mon-Sun columns

        # Prefer date from sheet name (row[0] is often stale in newer sheets)
        sheet_we = _sheet_date(name)

        for row in all_rows[1:]:
            emp = row[ci_name] if ci_name < len(row) else None
            if not emp or not isinstance(emp, str) or not emp.strip():
                continue
            emp = emp.strip()

            # Date: sheet name → row[0] → skip
            if sheet_we is not None:
                we = sheet_we
            elif isinstance(row[0], datetime):
                we = row[0]
            else:
                continue

            # Discard implausible dates (before 2020)
            if we.year < 2020:
                continue

            acct = str(row[ci_acct]).strip() if ci_acct < len(row) and row[ci_acct] else "–"
            sup  = str(row[ci_sup]).strip()  if ci_sup  < len(row) and row[ci_sup]  else "–"

            lates     = _safe_int(row[ci_late] if ci_late < len(row) else 0)
            left_early= _safe_int(row[ci_le]   if ci_le   < len(row) else 0)
            absent    = _safe_int(row[ci_abs]   if ci_abs  < len(row) else 0)
            sick      = _safe_int(row[ci_sick]  if ci_sick < len(row) else 0)
            leave     = _safe_int(row[ci_lv]    if ci_lv   < len(row) else 0)

            if ci_pres is not None and ci_pres < len(row):
                present = _safe_int(row[ci_pres])
            else:
                present = sum(
                    1 for c in day_cols
                    if c < len(row) and isinstance(row[c], str)
                    and row[c].strip().lower() == "present"
                )
                present += lates + left_early  # late/left-early = attended

            pct = None
            if ci_pct is not None and ci_pct < len(row):
                v = row[ci_pct]
                if isinstance(v, (int, float)):
                    pct = round(v * 100, 1) if v <= 1 else round(float(v), 1)
                elif isinstance(v, str) and v not in ("-", ""):
                    try:
                        pct = float(v.replace("%", ""))
                    except ValueError:
                        pass

            records.append({
                "week_ending": pd.Timestamp(we),
                "name":        emp,
                "account":     acct,
                "supervisor":  sup,
                "present":     present,
                "lates":       lates,
                "left_early":  left_early,
                "absent":      absent,
                "sick":        sick,
                "leave":       leave,
                "weekly_pct":  pct,
            })

    if not records:
        return pd.DataFrame()

    df = pd.DataFrame(records)
    df["week_ending"] = pd.to_datetime(df["week_ending"])
    df["year"]    = df["week_ending"].dt.year
    df["month"]   = df["week_ending"].dt.to_period("M")
    df["quarter"] = df["week_ending"].dt.to_period("Q")
    df["week"]    = df["week_ending"].dt.to_period("W")
    return df


# ── Worker thread ─────────────────────────────────────────────────────────────
class Loader(QThread):
    progress = pyqtSignal(int, int, str)
    done     = pyqtSignal(object)
    err      = pyqtSignal(str)

    def __init__(self, path):
        super().__init__()
        self.path = path

    def run(self):
        try:
            self.done.emit(load_excel(self.path, self.progress.emit))
        except Exception as e:
            self.err.emit(str(e))


# ── Chart ─────────────────────────────────────────────────────────────────────
class Chart(QWidget):
    def __init__(self):
        super().__init__()
        self.fig = Figure(figsize=(10, 3.8), facecolor=BG)
        self.canvas = FigureCanvas(self.fig)
        lay = QVBoxLayout(self)
        lay.setContentsMargins(0, 0, 0, 0)
        lay.addWidget(self.canvas)

    def plot(self, agg, period_label):
        self.fig.clear()
        ax = self.fig.add_subplot(111)
        ax.set_facecolor(BG_CARD)

        if agg.empty:
            ax.text(0.5, 0.5, "No data to display", color=TXT2,
                    ha="center", va="center", transform=ax.transAxes, fontsize=13)
            self.canvas.draw()
            return

        x      = range(len(agg))
        labels = [str(i) for i in agg.index]

        for col, color in STAT_COLORS.items():
            if col in agg.columns:
                ax.plot(x, agg[col], marker="o", label=col.replace("_", " ").title(),
                        color=color, linewidth=2, markersize=4)

        ax.set_xticks(list(x))
        ax.set_xticklabels(labels, rotation=45, ha="right", color=TXT2, fontsize=8)
        ax.tick_params(axis="y", colors=TXT2)
        for spine in ("top", "right"):
            ax.spines[spine].set_visible(False)
        ax.spines["bottom"].set_color(BORDER)
        ax.spines["left"].set_color(BORDER)
        ax.set_title(f"Attendance Trend  ·  by {period_label}", color=TXT,
                     fontsize=13, pad=10)
        ax.legend(framealpha=0.3, labelcolor=TXT, facecolor=BG_CARD,
                  edgecolor=BORDER, fontsize=9, loc="upper right")
        ax.grid(axis="y", color=BORDER, linestyle="--", alpha=0.5)
        self.fig.tight_layout()
        self.canvas.draw()


# ── Stat card ─────────────────────────────────────────────────────────────────
class Card(QFrame):
    def __init__(self, label, color):
        super().__init__()
        self.setStyleSheet(f"""
            QFrame {{
                background: {BG_CARD};
                border-left: 4px solid {color};
                border-radius: 8px;
            }}
        """)
        lay = QVBoxLayout(self)
        lay.setContentsMargins(14, 10, 14, 10)
        lay.setSpacing(2)

        self._lbl = QLabel(label)
        self._lbl.setStyleSheet(f"color:{TXT2}; font-size:11px; border:none;")
        self._val = QLabel("–")
        self._val.setStyleSheet(f"color:{color}; font-size:26px; font-weight:bold; border:none;")

        lay.addWidget(self._lbl)
        lay.addWidget(self._val)

    def set_value(self, v):
        self._val.setText(f"{v:,}" if isinstance(v, (int, float)) else str(v))


# ── Main window ───────────────────────────────────────────────────────────────
STYLE = f"""
QMainWindow, QWidget   {{ background:{BG};  color:{TXT}; font-family:'Segoe UI',Arial,sans-serif; }}
QLabel                 {{ color:{TXT}; }}
QPushButton            {{ background:{BLUE}; color:#fff; border:none; border-radius:6px; padding:8px 16px; font-size:13px; font-weight:bold; }}
QPushButton:hover      {{ background:#6ab0ff; }}
QPushButton#ghost      {{ background:transparent; color:{TXT2}; border:1px solid {BORDER}; }}
QPushButton#ghost:hover{{ color:{TXT}; border-color:{TXT2}; }}
QComboBox              {{ background:{BG_CARD}; color:{TXT}; border:1px solid {BORDER}; border-radius:6px; padding:6px 10px; font-size:12px; }}
QComboBox::drop-down   {{ border:none; width:20px; }}
QComboBox QAbstractItemView {{ background:{BG_CARD}; color:{TXT}; selection-background-color:{BLUE}; }}
QDateEdit              {{ background:{BG_CARD}; color:{TXT}; border:1px solid {BORDER}; border-radius:6px; padding:6px 10px; font-size:12px; }}
QDateEdit::drop-down   {{ border:none; width:20px; }}
QLineEdit              {{ background:{BG_CARD}; color:{TXT}; border:1px solid {BORDER}; border-radius:6px; padding:6px 10px; font-size:12px; }}
QGroupBox              {{ color:{TXT2}; border:1px solid {BORDER}; border-radius:8px; margin-top:10px; padding-top:6px; font-size:11px; font-weight:bold; }}
QGroupBox::title       {{ subcontrol-origin:margin; left:10px; top:-5px; color:{TXT2}; }}
QTabWidget::pane       {{ border:1px solid {BORDER}; border-radius:6px; }}
QTabBar::tab           {{ background:{BG_SIDE}; color:{TXT2}; padding:8px 22px; border-bottom:2px solid transparent; }}
QTabBar::tab:selected  {{ color:{BLUE}; border-bottom:2px solid {BLUE}; background:{BG}; }}
QTableWidget           {{ background:{BG_CARD}; color:{TXT}; gridline-color:{BORDER}; border:none; font-size:12px; alternate-background-color:#252538; }}
QTableWidget::item     {{ padding:5px; }}
QTableWidget::item:selected {{ background:{BLUE}; color:#fff; }}
QHeaderView::section   {{ background:{BG_SIDE}; color:{TXT2}; padding:8px; border:none; border-bottom:1px solid {BORDER}; font-weight:bold; font-size:11px; }}
QScrollBar:vertical    {{ background:{BG}; width:8px; border-radius:4px; }}
QScrollBar::handle:vertical {{ background:{BORDER}; border-radius:4px; }}
QSplitter::handle      {{ background:{BORDER}; }}
"""

PERIOD_MAP = {"Week": "week", "Month": "month", "Quarter": "quarter", "Year": "year"}


class App(QMainWindow):
    def __init__(self):
        super().__init__()
        self.df = None
        self.setWindowTitle("FSA Attendance Tracker")
        self.setMinimumSize(1280, 780)
        self.setStyleSheet(STYLE)
        self._build()

    # ── UI construction ───────────────────────────────────────────────────────
    def _build(self):
        root = QWidget()
        self.setCentralWidget(root)
        vlay = QVBoxLayout(root)
        vlay.setContentsMargins(0, 0, 0, 0)
        vlay.setSpacing(0)
        vlay.addWidget(self._topbar())

        split = QSplitter(Qt.Horizontal)
        split.setHandleWidth(1)
        split.addWidget(self._sidebar())
        split.addWidget(self._content())
        split.setSizes([240, 1040])
        vlay.addWidget(split)

    def _topbar(self):
        bar = QFrame()
        bar.setFixedHeight(56)
        bar.setStyleSheet(f"background:{BG_SIDE}; border-bottom:1px solid {BORDER};")
        lay = QHBoxLayout(bar)
        lay.setContentsMargins(20, 0, 20, 0)

        ttl = QLabel("FSA Attendance Tracker")
        ttl.setStyleSheet(f"font-size:18px; font-weight:bold; color:{TXT};")

        self._file_lbl = QLabel("No file loaded")
        self._file_lbl.setStyleSheet(f"color:{TXT2}; font-size:11px;")

        btn = QPushButton("Open Excel File")
        btn.setFixedWidth(150)
        btn.clicked.connect(self._open)

        lay.addWidget(ttl)
        lay.addSpacing(16)
        lay.addWidget(self._file_lbl)
        lay.addStretch()
        lay.addWidget(btn)
        return bar

    def _sidebar(self):
        side = QFrame()
        side.setFixedWidth(240)
        side.setStyleSheet(f"background:{BG_SIDE};")
        lay = QVBoxLayout(side)
        lay.setContentsMargins(12, 16, 12, 16)
        lay.setSpacing(12)

        # Period
        g1 = QGroupBox("View By")
        v1 = QVBoxLayout(g1)
        self._period = QComboBox()
        self._period.addItems(["Week", "Month", "Quarter", "Year"])
        self._period.setCurrentIndex(1)
        self._period.currentTextChanged.connect(self._refresh)
        v1.addWidget(self._period)
        lay.addWidget(g1)

        # Date range
        g2 = QGroupBox("Date Range")
        v2 = QVBoxLayout(g2)
        v2.addWidget(QLabel("From:"))
        self._d_from = QDateEdit(); self._d_from.setCalendarPopup(True)
        self._d_from.setDate(QDate(2023, 1, 1))
        v2.addWidget(self._d_from)
        v2.addWidget(QLabel("To:"))
        self._d_to = QDateEdit(); self._d_to.setCalendarPopup(True)
        self._d_to.setDate(QDate.currentDate())
        v2.addWidget(self._d_to)
        lay.addWidget(g2)

        # Filters
        g3 = QGroupBox("Filters")
        v3 = QVBoxLayout(g3)

        v3.addWidget(QLabel("Account:"))
        self._acct = QComboBox(); self._acct.addItem("All Accounts")
        v3.addWidget(self._acct)

        v3.addWidget(QLabel("Supervisor:"))
        self._sup = QComboBox(); self._sup.addItem("All Supervisors")
        v3.addWidget(self._sup)

        v3.addWidget(QLabel("Employee:"))
        self._emp = QLineEdit(); self._emp.setPlaceholderText("Search name…")
        v3.addWidget(self._emp)

        lay.addWidget(g3)

        apply_btn = QPushButton("Apply Filters")
        apply_btn.clicked.connect(self._refresh)
        lay.addWidget(apply_btn)

        reset_btn = QPushButton("Reset")
        reset_btn.setObjectName("ghost")
        reset_btn.clicked.connect(self._reset)
        lay.addWidget(reset_btn)

        lay.addStretch()
        return side

    def _content(self):
        w = QWidget()
        lay = QVBoxLayout(w)
        lay.setContentsMargins(16, 16, 16, 16)
        lay.setSpacing(12)
        self._tabs = QTabWidget()
        self._tabs.addTab(self._dash_tab(), "Dashboard")
        self._tabs.addTab(self._detail_tab(), "Details")
        lay.addWidget(self._tabs)
        return w

    def _dash_tab(self):
        w = QWidget()
        lay = QVBoxLayout(w)
        lay.setContentsMargins(0, 8, 0, 0)
        lay.setSpacing(12)

        # Stat cards
        row = QWidget()
        rlay = QHBoxLayout(row)
        rlay.setSpacing(10)
        rlay.setContentsMargins(0, 0, 0, 0)
        self._cards = {
            "present":    Card("Total Present",    GREEN),
            "absent":     Card("Total Absent",     RED),
            "lates":      Card("Total Lates",      YELLOW),
            "left_early": Card("Total Left Early", ORANGE),
            "sick":       Card("Total Sick",       PURPLE),
            "leave":      Card("Total Leave",      TEAL),
        }
        for c in self._cards.values():
            rlay.addWidget(c)
        lay.addWidget(row)

        self._chart = Chart()
        lay.addWidget(self._chart)
        return w

    def _detail_tab(self):
        w = QWidget()
        lay = QVBoxLayout(w)
        lay.setContentsMargins(0, 8, 0, 0)
        self._table = QTableWidget()
        self._table.setAlternatingRowColors(True)
        self._table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeToContents)
        self._table.horizontalHeader().setStretchLastSection(True)
        self._table.verticalHeader().setVisible(False)
        self._table.setSortingEnabled(True)
        lay.addWidget(self._table)
        return w

    # ── File loading ──────────────────────────────────────────────────────────
    def _open(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "Open Excel File", "", "Excel Files (*.xlsx *.xls)")
        if not path:
            return

        self._prog = QProgressDialog("Loading…", None, 0, 100, self)
        self._prog.setWindowTitle("Loading")
        self._prog.setWindowModality(Qt.WindowModal)
        self._prog.setMinimumDuration(0)
        self._prog.setValue(0)

        self._worker = Loader(path)
        self._worker.progress.connect(self._on_prog)
        self._worker.done.connect(self._on_done)
        self._worker.err.connect(self._on_err)
        self._worker.start()

        self._file_lbl.setText(os.path.basename(path))

    def _on_prog(self, cur, total, name):
        self._prog.setValue(int(cur / total * 100) if total else 0)
        self._prog.setLabelText(f"Reading: {name}")

    def _on_done(self, df):
        self._prog.close()
        if df.empty:
            QMessageBox.warning(self, "Warning", "No attendance data found.")
            return
        self.df = df
        self._fill_combos()
        self._refresh()

    def _on_err(self, msg):
        self._prog.close()
        QMessageBox.critical(self, "Error", f"Failed to load:\n{msg}")

    # ── Filters ───────────────────────────────────────────────────────────────
    def _fill_combos(self):
        accts = sorted(self.df["account"].dropna().unique())
        self._acct.clear(); self._acct.addItem("All Accounts"); self._acct.addItems(accts)

        sups = sorted(self.df["supervisor"].dropna().unique())
        self._sup.clear(); self._sup.addItem("All Supervisors"); self._sup.addItems(sups)

        lo = self.df["week_ending"].min()
        hi = self.df["week_ending"].max()
        self._d_from.setDate(QDate(lo.year, lo.month, lo.day))
        self._d_to.setDate(QDate(hi.year, hi.month, hi.day))

    def _reset(self):
        self._period.setCurrentIndex(1)
        self._emp.clear()
        if self.df is not None:
            self._fill_combos()
            self._refresh()

    def _filtered(self):
        if self.df is None:
            return None
        df = self.df.copy()

        d0 = self._d_from.date()
        d1 = self._d_to.date()
        t0 = pd.Timestamp(d0.year(), d0.month(), d0.day())
        t1 = pd.Timestamp(d1.year(), d1.month(), d1.day())
        df = df[(df["week_ending"] >= t0) & (df["week_ending"] <= t1)]

        if self._acct.currentIndex() > 0:
            df = df[df["account"] == self._acct.currentText()]
        if self._sup.currentIndex() > 0:
            df = df[df["supervisor"] == self._sup.currentText()]

        srch = self._emp.text().strip()
        if srch:
            df = df[df["name"].str.contains(srch, case=False, na=False)]

        return df

    # ── Refresh ───────────────────────────────────────────────────────────────
    def _refresh(self):
        df = self._filtered()
        if df is None or df.empty:
            return
        self._update_dash(df)
        self._update_table(df)

    def _period_col(self):
        return PERIOD_MAP[self._period.currentText()]

    def _update_dash(self, df):
        for key, card in self._cards.items():
            card.set_value(df[key].sum())

        pcol = self._period_col()
        agg  = df.groupby(pcol)[list(STAT_COLORS)].sum()
        # Limit chart to last 52 data points so it stays readable
        if len(agg) > 52:
            agg = agg.tail(52)
        self._chart.plot(agg, self._period.currentText())

    def _update_table(self, df):
        pcol    = self._period_col()
        stat_cols = list(STAT_COLORS)
        display_cols = ["name", "account", "supervisor", pcol] + stat_cols + ["weekly_pct"]
        display_cols = [c for c in display_cols if c in df.columns]

        view = df[display_cols].copy()
        view[pcol] = view[pcol].astype(str)
        view = view.sort_values([pcol, "name"], ascending=[False, True])

        headers = {
            "name":        "Employee",
            "account":     "Account",
            "supervisor":  "Supervisor",
            pcol:          self._period.currentText(),
            "present":     "Present",
            "absent":      "Absent",
            "lates":       "Lates",
            "left_early":  "Left Early",
            "sick":        "Sick",
            "leave":       "Leave",
            "weekly_pct":  "Attendance %",
        }

        self._table.setSortingEnabled(False)
        self._table.clear()
        self._table.setRowCount(len(view))
        self._table.setColumnCount(len(display_cols))
        self._table.setHorizontalHeaderLabels([headers.get(c, c) for c in display_cols])

        color_map = {k: QColor(v) for k, v in STAT_COLORS.items()}

        for r, (_, row) in enumerate(view.iterrows()):
            for c, col in enumerate(display_cols):
                val = row[col]
                text = "" if val is None else \
                       f"{val:.1f}%" if col == "weekly_pct" and isinstance(val, float) else \
                       str(val)
                item = QTableWidgetItem(text)
                item.setFlags(Qt.ItemIsSelectable | Qt.ItemIsEnabled)
                if col in color_map and isinstance(val, (int, float)) and val > 0:
                    item.setForeground(color_map[col])
                # Right-align numeric columns
                if col in {*STAT_COLORS, "weekly_pct"}:
                    item.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
                self._table.setItem(r, c, item)

        self._table.setSortingEnabled(True)


# ── Entry point ───────────────────────────────────────────────────────────────
if __name__ == "__main__":
    app = QApplication(sys.argv)
    app.setStyle("Fusion")
    win = App()
    win.show()
    sys.exit(app.exec_())
